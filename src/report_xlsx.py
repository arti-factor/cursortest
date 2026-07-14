"""Erzeugt den Excel-Report (Übersicht, Kriteriengruppen, Maßnahmenliste, ...).

Funktioniert identisch bei einem wie bei mehreren hundert Standorten - es gibt
keine standortzahl-abhängige Sonderlogik.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.checks import GROUPS
from src.models import ClientRun, LocationAuditResult, Modus
from src.scoring import top_massnahmen

_AMPEL_FILL = {
    "gruen": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "gelb": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "rot": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "n/a": PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"),
}
_STATUS_FILL = {
    "ok": _AMPEL_FILL["gruen"],
    "warnung": _AMPEL_FILL["gelb"],
    "fehler": _AMPEL_FILL["rot"],
    "nicht_pruefbar": _AMPEL_FILL["n/a"],
}
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

_GRUPPEN_LABELS = {
    "existenz_auffindbarkeit": "Existenz & Auffindbarkeit",
    "basisdaten_nap": "Basisdaten & NAP",
    "kategorien": "Kategorien",
    "oeffnungszeiten": "Öffnungszeiten",
    "beschreibung_attribute": "Beschreibung & Attribute",
    "fotos": "Fotos",
    "bewertungen": "Bewertungen",
    "qna": "Q&A",
    "leistungen": "Leistungen und Produkte",
    "performance": "Performance-Baseline",
}

_MODUS_B_GRUPPEN = ("beschreibung_attribute", "qna", "leistungen", "performance")


_INVALID_SHEET_CHARS = str.maketrans("", "", "[]:*?/\\")


def _safe_sheet_name(name: str) -> str:
    return name.translate(_INVALID_SHEET_CHARS)[:31]


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_uebersicht(wb: Workbook, results: list[LocationAuditResult], top_n: int) -> None:
    ws = wb.active
    ws.title = "Übersicht"
    gruppen_keys = [g for g, _ in GROUPS]
    headers = (
        ["Standort", "Adresse", "Modus", "Gesamtscore", "Ampel"]
        + [_GRUPPEN_LABELS[g] for g in gruppen_keys]
        + [f"Mangel {i}" for i in range(1, top_n + 1)]
    )
    _write_header(ws, headers)

    for row_idx, result in enumerate(results, start=2):
        ws.cell(row=row_idx, column=1, value=result.location.name)
        ws.cell(row=row_idx, column=2, value=result.location.address)
        ws.cell(row=row_idx, column=3, value=result.modus.value)
        ws.cell(row=row_idx, column=4, value=result.gesamtscore)
        ampel_cell = ws.cell(row=row_idx, column=5, value=result.ampel)
        ampel_cell.fill = _AMPEL_FILL.get(result.ampel, _AMPEL_FILL["n/a"])

        gruppen_by_key = {gs.gruppe: gs for gs in result.gruppen}
        for i, g in enumerate(gruppen_keys):
            col = 6 + i
            gs = gruppen_by_key.get(g)
            cell = ws.cell(row=row_idx, column=col, value=gs.punkte if gs else None)
            if gs and gs.lauffaehig and gs.punkte is not None:
                if gs.punkte >= 85:
                    cell.fill = _AMPEL_FILL["gruen"]
                elif gs.punkte >= 60:
                    cell.fill = _AMPEL_FILL["gelb"]
                else:
                    cell.fill = _AMPEL_FILL["rot"]
            elif gs and not gs.lauffaehig:
                cell.fill = _AMPEL_FILL["n/a"]

        mangel_start_col = 6 + len(gruppen_keys)
        top = top_massnahmen(result, anzahl=top_n)
        for i in range(top_n):
            wert = top[i].befund if i < len(top) else ""
            ws.cell(row=row_idx, column=mangel_start_col + i, value=wert)

    _autosize(ws, [28, 36, 8, 12, 10] + [14] * len(gruppen_keys) + [50] * top_n)


def _write_gruppen_blaetter(wb: Workbook, results: list[LocationAuditResult]) -> None:
    for gruppe_key, _ in GROUPS:
        ws = wb.create_sheet(_safe_sheet_name(_GRUPPEN_LABELS[gruppe_key]))
        headers = ["Standort", "Check", "Status", "Punkte", "Befund", "Empfehlung"]
        _write_header(ws, headers)
        row_idx = 2
        for result in results:
            gs = next((g for g in result.gruppen if g.gruppe == gruppe_key), None)
            if not gs:
                continue
            for check in gs.checks:
                ws.cell(row=row_idx, column=1, value=result.location.name)
                ws.cell(row=row_idx, column=2, value=check.check_id)
                status_cell = ws.cell(row=row_idx, column=3, value=check.status.value)
                status_cell.fill = _STATUS_FILL.get(check.status.value, _AMPEL_FILL["n/a"])
                ws.cell(row=row_idx, column=4, value=check.punkte)
                ws.cell(row=row_idx, column=5, value=check.befund)
                ws.cell(row=row_idx, column=6, value=check.empfehlung)
                row_idx += 1
        _autosize(ws, [28, 30, 14, 10, 60, 50])


def _write_massnahmenliste(wb: Workbook, results: list[LocationAuditResult]) -> None:
    ws = wb.create_sheet("Maßnahmenliste")
    headers = ["Standort", "Kriterium", "Status", "Befund", "Empfohlene Maßnahme", "Impact"]
    _write_header(ws, headers)

    zeilen = []
    for result in results:
        gewicht_je_gruppe = {gs.gruppe: gs.gewicht_prozent for gs in result.gruppen}
        for gs in result.gruppen:
            for check in gs.checks:
                if check.status.value not in ("warnung", "fehler"):
                    continue
                abzug = 100 - (check.punkte if check.punkte is not None else 0)
                impact = gewicht_je_gruppe.get(check.gruppe, 0.0) * abzug
                zeilen.append((result.location.name, check, impact))

    zeilen.sort(key=lambda z: z[2], reverse=True)
    for row_idx, (standort_name, check, impact) in enumerate(zeilen, start=2):
        ws.cell(row=row_idx, column=1, value=standort_name)
        ws.cell(row=row_idx, column=2, value=check.check_id)
        status_cell = ws.cell(row=row_idx, column=3, value=check.status.value)
        status_cell.fill = _STATUS_FILL.get(check.status.value, _AMPEL_FILL["n/a"])
        ws.cell(row=row_idx, column=4, value=check.befund)
        ws.cell(row=row_idx, column=5, value=check.empfehlung)
        ws.cell(row=row_idx, column=6, value=round(impact, 1))

    _autosize(ws, [28, 32, 14, 60, 50, 10])


def _write_nicht_geprueft(wb: Workbook, results: list[LocationAuditResult]) -> None:
    ws = wb.create_sheet("Nicht geprüft")
    headers = ["Standort", "Kriteriengruppe", "Check", "Hinweis"]
    _write_header(ws, headers)
    row_idx = 2
    for result in results:
        for gs in result.gruppen:
            if gs.gruppe not in _MODUS_B_GRUPPEN and gs.lauffaehig:
                continue
            for check in gs.checks:
                if check.status.value != "nicht_pruefbar":
                    continue
                ws.cell(row=row_idx, column=1, value=result.location.name)
                ws.cell(row=row_idx, column=2, value=_GRUPPEN_LABELS.get(gs.gruppe, gs.gruppe))
                ws.cell(row=row_idx, column=3, value=check.check_id)
                ws.cell(row=row_idx, column=4, value=check.befund)
                row_idx += 1
    ws.cell(
        row=row_idx + 1,
        column=1,
        value=(
            "Diese Prüfungen sind nur mit Verwaltungszugriff auf das Google-Unternehmensprofil "
            "möglich (Full Audit). Siehe README.md, Abschnitt 'Modus B'."
        ),
    )
    _autosize(ws, [28, 28, 32, 70])


def _write_performance(wb: Workbook, results: list[LocationAuditResult]) -> None:
    ws = wb.create_sheet("Performance")
    headers = ["Standort", "Check", "Status", "Befund"]
    _write_header(ws, headers)
    row_idx = 2
    for result in results:
        gs = next((g for g in result.gruppen if g.gruppe == "performance"), None)
        if not gs:
            continue
        for check in gs.checks:
            if check.status.value == "nicht_pruefbar":
                continue
            ws.cell(row=row_idx, column=1, value=result.location.name)
            ws.cell(row=row_idx, column=2, value=check.check_id)
            status_cell = ws.cell(row=row_idx, column=3, value=check.status.value)
            status_cell.fill = _STATUS_FILL.get(check.status.value, _AMPEL_FILL["n/a"])
            ws.cell(row=row_idx, column=4, value=check.befund)
            row_idx += 1
    _autosize(ws, [28, 24, 14, 80])


def write_report(client_run: ClientRun, output_path: Path, report_config: dict[str, Any]) -> Path:
    wb = Workbook()
    top_n = report_config.get("top_massnahmen_je_standort", 3)

    _write_uebersicht(wb, client_run.results, top_n)
    _write_gruppen_blaetter(wb, client_run.results)
    _write_massnahmenliste(wb, client_run.results)

    if client_run.modus == Modus.A:
        _write_nicht_geprueft(wb, client_run.results)
    else:
        _write_performance(wb, client_run.results)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
