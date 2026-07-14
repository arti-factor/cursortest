from pathlib import Path

from openpyxl import load_workbook

from src.checks.base import fehler, ok
from src.models import ClientRun, GbpProfile, Location, Modus
from src.report_html import write_html_report
from src.report_xlsx import write_report
from src.scoring import score_location
from src.storage import load_global_config


def _make_result(loc_id: str, name: str, modus=Modus.A):
    location = Location(id=loc_id, name=name, zip="12345", city="Musterstadt")
    profile = GbpProfile(place_id=f"ChIJ_{loc_id}", display_name=name)
    checks = [
        ok("existenz_gbp_vorhanden", "existenz_auffindbarkeit", modus, "Profil gefunden", punkte=100.0),
        fehler("nap_telefon_vorhanden", "basisdaten_nap", modus, "Telefon fehlt", "ergänzen", punkte=0.0),
    ]
    return score_location(location, profile, modus, checks, load_global_config()["scoring"])


def test_write_report_erzeugt_alle_erwarteten_blaetter(tmp_path: Path):
    results = [_make_result("a", "Standort A"), _make_result("b", "Standort B")]
    run = ClientRun(client_slug="testkunde", run_date="2026-07-14", modus=Modus.A, results=results)

    output_path = tmp_path / "report.xlsx"
    write_report(run, output_path, load_global_config()["report"])

    assert output_path.exists()
    wb = load_workbook(output_path)
    assert "Übersicht" in wb.sheetnames
    assert "Maßnahmenliste" in wb.sheetnames
    assert "Nicht geprüft" in wb.sheetnames
    assert "Performance" not in wb.sheetnames  # nur im Modus B

    uebersicht = wb["Übersicht"]
    assert uebersicht.cell(row=2, column=1).value == "Standort A"
    assert uebersicht.cell(row=3, column=1).value == "Standort B"


def test_write_report_funktioniert_mit_einem_einzigen_standort(tmp_path: Path):
    results = [_make_result("a", "Einzelstandort")]
    run = ClientRun(client_slug="einzelkunde", run_date="2026-07-14", modus=Modus.B, results=results)

    output_path = tmp_path / "report_einzel.xlsx"
    write_report(run, output_path, load_global_config()["report"])

    wb = load_workbook(output_path)
    assert "Performance" in wb.sheetnames
    assert "Nicht geprüft" not in wb.sheetnames
    assert wb["Übersicht"].max_row == 2  # Header + 1 Standort


def test_write_html_report(tmp_path: Path):
    results = [_make_result("a", "Standort A")]
    run = ClientRun(client_slug="testkunde", run_date="2026-07-14", modus=Modus.A, results=results)

    output_path = tmp_path / "dashboard.html"
    write_html_report(run, output_path)

    content = output_path.read_text(encoding="utf-8")
    assert "Standort A" in content
    assert "testkunde" in content
    assert "<table" in content
